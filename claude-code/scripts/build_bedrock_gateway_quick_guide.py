from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"deliverables\Prisma_AIRS_AI_Gateway_Bedrock_Quick_Guide.docx"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def borders(table, color="D9D9D9"):
    tbl_pr = table._tbl.tblPr
    element = tbl_pr.find(qn("w:tblBorders"))
    if element is None:
        element = OxmlElement("w:tblBorders")
        tbl_pr.append(element)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = OxmlElement(f"w:{edge}")
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), "4")
        tag.set(qn("w:color"), color)
        element.append(tag)


def set_cell_text(cell, text, bold=False, color="000000"):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    run.bold = bold
    run.font.name = "Aptos"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_numbered(doc, number, heading, body):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(f"{number}. {heading}")
    r.bold = True
    r.font.name = "Aptos"
    r.font.size = Pt(10.5)
    p2 = doc.add_paragraph(body)
    p2.paragraph_format.left_indent = Inches(0.22)
    p2.paragraph_format.space_after = Pt(4)


doc = Document()
section = doc.sections[0]
section.top_margin = Inches(0.45)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)

styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(9.5)
styles["Normal"].font.color.rgb = RGBColor(0, 0, 0)
styles["Normal"].paragraph_format.space_after = Pt(4)

# Remove any inherited border from Word's built-in Title style.
title_ppr = styles["Title"]._element.get_or_add_pPr()
title_borders = title_ppr.find(qn("w:pBdr"))
if title_borders is not None:
    title_ppr.remove(title_borders)

title = doc.add_paragraph(style="Title")
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after = Pt(3)
r = title.add_run("Prisma AIRS AI Gateway Amazon Bedrock Quick Guide")
r.font.name = "Aptos Display"
r.font.size = Pt(20)
r.font.color.rgb = RGBColor(0, 0, 0)

sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
sub.paragraph_format.space_after = Pt(8)
r = sub.add_run("The minimum AWS configuration for the Bedrock integration screen")
r.italic = True
r.font.size = Pt(9.5)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(7)
r = p.add_run("Recommended for this screen. ")
r.bold = True
p.add_run("Select AWS Access Key. Create a dedicated IAM user for AIRS with Bedrock invocation permissions only. Do not select API Key for the native Amazon Bedrock integration.")

add_numbered(doc, 1, "Enable the Bedrock model", "In the required AWS Region, confirm that the approved model is available and can be invoked from the Amazon Bedrock playground. Complete any one-time provider or Marketplace activation using an administrator identity, not the AIRS integration identity.")

add_numbered(doc, 2, "Create the IAM policy", "In AWS IAM, create a customer-managed policy named PrismaAIRSBedrockInvoke and paste the policy below. This is the simple initial policy. After validation, restrict Resource to the approved model and inference-profile ARNs.")

policy = '''{
  "Version": "2012-10-17",
  "Statement": [{
    "Effect": "Allow",
    "Action": [
      "bedrock:InvokeModel",
      "bedrock:InvokeModelWithResponseStream",
      "bedrock:GetInferenceProfile"
    ],
    "Resource": "*"
  }]
}'''
p = doc.add_paragraph()
p.paragraph_format.left_indent = Inches(0.22)
p.paragraph_format.right_indent = Inches(0.22)
p.paragraph_format.space_after = Pt(6)
for i, line in enumerate(policy.splitlines()):
    rr = p.add_run(line + ("\n" if i < len(policy.splitlines()) - 1 else ""))
    rr.font.name = "Consolas"
    rr.font.size = Pt(8)

add_numbered(doc, 3, "Create the integration identity", "Create a dedicated IAM user such as prisma-airs-bedrock, attach only PrismaAIRSBedrockInvoke, and create one access key. Copy the Access Key ID and Secret Access Key once. Store the secret in the approved password or secrets manager and never send it by email.")

add_numbered(doc, 4, "Complete the AIRS Integration Details", "Enter the generated AWS values in the Bedrock integration form as follows.")

table = doc.add_table(rows=1, cols=2)
table.autofit = False
table.columns[0].width = Inches(2.35)
table.columns[1].width = Inches(4.65)
headers = table.rows[0].cells
set_cell_text(headers[0], "AIRS field", True, "FFFFFF")
set_cell_text(headers[1], "Value", True, "FFFFFF")
shade(headers[0], "243447")
shade(headers[1], "243447")
rows = [
    ("Name", "A descriptive name, for example Customer Bedrock"),
    ("Slug", "A unique lowercase identifier, for example customer-bedrock"),
    ("AWS Auth Type", "AWS Access Key"),
    ("AWS Secret Access Key", "Secret from the dedicated IAM access key"),
    ("AWS Access Key Id", "Access Key ID from the same IAM access key"),
    ("AWS Region", "The Region containing the approved Bedrock model, for example us-east-1"),
    ("Custom Host and Headers", "Leave blank unless the customer uses a documented private endpoint or proxy"),
]
for idx, (field, value) in enumerate(rows):
    cells = table.add_row().cells
    set_cell_text(cells[0], field, True)
    set_cell_text(cells[1], value)
    if idx % 2:
        shade(cells[0], "F3F6F9")
        shade(cells[1], "F3F6F9")
borders(table)

add_numbered(doc, 5, "Provision and verify", "Provision only the required workspace and approved model. Create the integration, run the connection test, and send one harmless test prompt through AI Gateway. Confirm the request appears in AI Gateway observability and succeeds in Bedrock.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(3)
r = p.add_run("Security note")
r.bold = True
r.font.size = Pt(10.5)
doc.add_paragraph("AWS Assumed Role is preferable when the customer prohibits long-lived access keys. Use that option only when the AIRS form provides the exact trusted principal and External ID required for the role trust policy. Do not invent either value. Rotate or revoke the access key after the POC and do not attach AdministratorAccess or AmazonBedrockFullAccess.")

p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(5)
p.paragraph_format.space_after = Pt(2)
r = p.add_run("Official references")
r.bold = True
r.font.size = Pt(10.5)
for text in [
    "Palo Alto Networks: Configure AI Gateway — https://docs.paloaltonetworks.com/ai-runtime-security/administration/configure-ai-gateway",
    "AWS: Bedrock inference permissions — https://docs.aws.amazon.com/bedrock/latest/userguide/inference-prereq.html",
    "AWS: Model access — https://docs.aws.amazon.com/bedrock/latest/userguide/model-access.html",
]:
    p = doc.add_paragraph(text)
    p.paragraph_format.left_indent = Inches(0.1)
    p.paragraph_format.space_after = Pt(0)
    for run in p.runs:
        run.font.size = Pt(7.25)

doc.save(OUTPUT)
print(OUTPUT)
