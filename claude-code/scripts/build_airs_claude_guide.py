from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MEDIA = ROOT / "outputs" / "guide-inspect" / "expanded" / "word" / "media"
OUT = ROOT / "outputs" / "Prisma_AIRS_AI_Gateway_Claude_Code_Guide.docx"


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar")
        tcPr.append(tcMar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tcMar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tcMar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(9)
    r.font.color.rgb = RGBColor(89, 89, 89)


def add_image(doc, filename, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(MEDIA / filename), width=Inches(6.75))
    add_caption(doc, caption)


def add_numbered_step(doc, number, heading, body):
    p = doc.add_paragraph()
    p.style = doc.styles["Heading 2"]
    p.add_run(f"{number}. {heading}")
    b = doc.add_paragraph(body)
    b.paragraph_format.space_after = Pt(8)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7)
sec.bottom_margin = Inches(0.7)
sec.left_margin = Inches(0.8)
sec.right_margin = Inches(0.8)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor(32, 33, 36)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.08

title = styles["Title"]
title.font.name = "Aptos Display"
title.font.size = Pt(25)
title.font.bold = True
title.font.color.rgb = RGBColor(0, 0, 0)
title.paragraph_format.space_after = Pt(6)

for style_name, size in (("Heading 1", 17), ("Heading 2", 13)):
    style = styles[style_name]
    style.font.name = "Aptos Display"
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(13 if style_name == "Heading 1" else 9)
    style.paragraph_format.space_after = Pt(5)

if "Code" not in styles:
    code_style = styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
else:
    code_style = styles["Code"]
code_style.font.name = "Cascadia Mono"
code_style.font.size = Pt(9)
code_style.font.color.rgb = RGBColor(32, 33, 36)
code_style.paragraph_format.left_indent = Inches(0.25)
code_style.paragraph_format.right_indent = Inches(0.25)
code_style.paragraph_format.space_before = Pt(5)
code_style.paragraph_format.space_after = Pt(8)

p = doc.add_paragraph(style="Title")
p.add_run("Prisma AIRS AI Gateway for Claude Code")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
r = p.add_run("Windows setup using a Gateway service key and an attached routing configuration")
r.bold = True
r.font.size = Pt(12)
r.font.color.rgb = RGBColor(52, 84, 143)

doc.add_paragraph(
    "This guide configures Claude Code to send model requests through Prisma AIRS AI Gateway. "
    "The Gateway key authenticates Claude Code, while the attached routing configuration selects "
    "the provider route and applies the configuration's routing and guardrail hooks."
)

doc.add_heading("Applies to", level=1)
t = doc.add_table(rows=4, cols=2)
t.style = "Table Grid"
rows = [
    ("Management plane", "Strata Cloud Manager"),
    ("Gateway", "Prisma AIRS AI Gateway SaaS"),
    ("Client", "Claude Code on Windows"),
    ("Authentication", "Gateway service key stored securely on the endpoint"),
]
for i, (a, b) in enumerate(rows):
    t.cell(i, 0).text = a
    t.cell(i, 1).text = b
    for cell in t.rows[i].cells:
        set_cell_margins(cell)
    t.cell(i, 0).paragraphs[0].runs[0].bold = True

doc.add_heading("Prerequisites", level=1)
for item in [
    "Prisma AIRS AI Gateway is activated in the Americas region.",
    "At least one LLM integration is configured and provisioned to the target workspace.",
    "A saved Gateway routing configuration exists for the intended provider and guardrails.",
    "You have Workspace Manager or Organization Admin access in Strata Cloud Manager.",
    "Claude Code is installed on the Windows endpoint.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Traffic flow", level=1)
doc.add_paragraph(
    "Claude Code sends an Anthropic-format request to the AI Gateway base URL. The Gateway key "
    "authenticates the caller. The workspace virtual key identifies the workspace when required. "
    "The routing configuration attached to the Gateway key is then applied automatically before "
    "the request is forwarded to the configured LLM provider."
)

doc.add_heading("Configure Prisma AIRS AI Gateway", level=1)
add_numbered_step(
    doc,
    1,
    "Review organization security settings",
    "In Strata Cloud Manager, open AI Security > AI Gateway > Admin Settings > Security. Review the "
    "inline-configuration controls. If you want the saved key configuration to be enforced, enable "
    "Block Inline Configurations so clients cannot replace it with inline JSON.",
)
add_image(doc, "image3.png", "Organization-level controls for inline configurations")

add_numbered_step(
    doc,
    2,
    "Open Security Keys",
    "Go to AI Security > AI Gateway > Security Keys. Select Service for a backend or centrally "
    "managed Claude Code deployment, and then select Create New.",
)
add_image(doc, "image2.png", "Security Keys page with the Service tab and Create New action")

add_numbered_step(
    doc,
    3,
    "Create the Gateway API key",
    "Enter a descriptive name. In Configuration, select the saved Gateway routing configuration. "
    "When a configuration is attached to the key, it is applied automatically to every request that "
    "uses that key. The client does not need to send x-portkey-config. Disable configuration override "
    "when the client must not select a different configuration.",
)
add_image(doc, "image4.png", "Gateway service key with a saved routing configuration attached")

add_numbered_step(
    doc,
    4,
    "Assign permissions",
    "Grant only the permissions required by the deployment. For basic Claude Code inference, confirm "
    "that the key can invoke the required completion path. Add Agents, Logs, MCP, or Prompt permissions "
    "only when those capabilities are part of the intended workflow. Select Create Gateway API Key.",
)
add_image(doc, "image1.png", "Permission selection before creating the Gateway API key")

add_numbered_step(
    doc,
    5,
    "Copy and protect the key",
    "Copy the key when it is displayed. The value is shown once and cannot be retrieved after the "
    "dialog is closed. Store it in an approved secret store. Do not commit it to Git or place it in "
    "a shared project settings file.",
)

doc.add_heading("Configure Claude Code on Windows", level=1)
add_numbered_step(
    doc,
    6,
    "Open the user settings file",
    r"Use %USERPROFILE%\.claude\settings.json. This file applies to all projects for the current Windows user.",
)

add_numbered_step(
    doc,
    7,
    "Add the Gateway settings",
    "Use the Gateway URL shown in AI Gateway Admin Settings. Replace the example values below. If your "
    "workspace requires a virtual key, include it as shown. The attached routing configuration means "
    "x-portkey-config is not required.",
)

code = r'''{
  "env": {
    "ANTHROPIC_BASE_URL": "YOUR_AIRS_GATEWAY_URL",
    "ANTHROPIC_AUTH_TOKEN": "YOUR_GATEWAY_API_KEY",
    "ANTHROPIC_CUSTOM_HEADERS": "x-portkey-api-key: YOUR_GATEWAY_API_KEY\nx-portkey-virtual-key: YOUR_WORKSPACE_VIRTUAL_KEY\nx-portkey-provider: YOUR_PROVIDER_SLUG"
  }
}'''
p = doc.add_paragraph(style="Code")
p.add_run(code)

p = doc.add_paragraph()
r = p.add_run("Security note: ")
r.bold = True
p.add_run(
    "For an organization-wide rollout, distribute the base URL through managed settings and deliver "
    "each user's credential through approved secrets tooling or apiKeyHelper. Avoid publishing a shared "
    "Gateway key in a common file."
)

doc.add_heading("Verify", level=1)
for item in [
    "Start Claude Code and run /status.",
    "Confirm Anthropic base URL shows the Prisma AIRS AI Gateway URL.",
    "Confirm the active credential source is the configured token, API key, or apiKeyHelper rather than the saved Claude subscription login.",
    "Send a distinctive test prompt and confirm that it returns normally.",
    "In AI Gateway Observability, confirm the request appears under the expected workspace, provider, and attached configuration.",
]:
    doc.add_paragraph(item, style="List Number")

doc.add_heading("Important behavior", level=1)
for item in [
    "A Gateway credential replaces the Claude subscription credential for those requests. The saved Claude login remains on the machine but is not sent while the Gateway credential is active.",
    "Attaching a routing configuration to the key removes the need to send x-portkey-config on every request.",
    "The pc-prefixed value identifies an AI Gateway routing configuration. It is not an AIRS Runtime API security-profile UUID.",
    "Claude Desktop uses a separate third-party inference configuration. It does not use Claude Code's settings.json for Gateway routing.",
]:
    doc.add_paragraph(item, style="List Bullet")

doc.add_heading("Rollback", level=1)
doc.add_paragraph(
    "Remove the Gateway variables from %USERPROFILE%\\.claude\\settings.json and restart Claude Code. "
    "Claude Code then returns to its normal saved-login behavior. Revoke the Gateway key in Strata Cloud "
    "Manager if it is no longer required."
)

doc.add_heading("Official references", level=1)
refs = [
    "Palo Alto Networks — Manage Gateway API Keys: https://docs.paloaltonetworks.com/prisma-airs/ai-gateway/ai-gateway-configs/manage-gateway-api-keys",
    "Palo Alto Networks — AI Gateway Routing Configurations: https://docs.paloaltonetworks.com/prisma-airs/ai-gateway/ai-gateway-configs",
    "Palo Alto Networks — AI Gateway Security Keys: https://docs.paloaltonetworks.com/prisma-airs/ai-gateway/ai-gateway-security-keys",
    "Anthropic — Connect Claude Code to an LLM gateway: https://code.claude.com/docs/en/llm-gateway-connect",
    "Anthropic — Other LLM gateways: https://code.claude.com/docs/en/llm-gateway",
]
for ref in refs:
    doc.add_paragraph(ref, style="List Bullet")

for section in doc.sections:
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Prisma AIRS AI Gateway for Claude Code")
    run.font.name = "Aptos"
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(120, 120, 120)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
